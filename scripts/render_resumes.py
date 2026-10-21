"""Create editable documents from local private TXT drafts. No embedded personal data."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK
import argparse
ap=argparse.ArgumentParser();ap.add_argument('--job');args=ap.parse_args()
root=Path(__file__).resolve().parents[1]
for folder in sorted((root/'private/jobs').iterdir()):
 if args.job and folder.name!=args.job:continue
 if not folder.is_dir():continue
 for src in folder.glob('03*.txt'):
  doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.72);s.right_margin=Inches(.72)
  normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0);normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.05
  for name in ['Title','Subtitle','Heading 1','Heading 2']:
   st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
  doc.styles['Title'].font.size=Pt(23);doc.styles['Title'].paragraph_format.space_after=Pt(5)
  doc.styles['Heading 1'].font.size=Pt(11.5);doc.styles['Heading 1'].font.bold=True;doc.styles['Heading 1'].paragraph_format.space_before=Pt(9);doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
  from docx.oxml.ns import qn
  for elem in list(doc.styles.element.iter(qn('w:pBdr'))): elem.getparent().remove(elem)
  lines=src.read_text().splitlines();first=True;section_name='';education_paras=[]
  for line in lines:
   line=line.strip().replace('–','-').replace('—','-').replace('‑','-')
   if not line:continue
   if line.startswith('#'):line=line.lstrip('# ').strip()
   if first:doc.add_paragraph(line,style='Title');first=False;continue
   if line.isupper() and len(line)<75 and not line.startswith('-'):
    section_name=line
    para=doc.add_paragraph(line.title(),style='Heading 1')
    if line=='EDUCATION':education_paras.append(para)
    continue
   if line.startswith(('- ','• ')):
    para=doc.add_paragraph(style='List Bullet');para.add_run(line[2:]);para.paragraph_format.space_after=Pt(4)
   else:
    para=doc.add_paragraph(line)
    if ' | ' in line or 'Research Intern' in line or 'Visiting Student Researcher' in line or line.startswith(('Appointment commencing','Research Intern','PhD research','BSc ','B.Sc.','PhD in','University of')):
     para.runs[0].bold=True;para.paragraph_format.keep_with_next=True
   para.paragraph_format.widow_control=True
   if section_name=='EDUCATION':education_paras.append(para)
  for idx,para in enumerate(doc.paragraphs[:-1]):
   nxt=doc.paragraphs[idx+1]
   if nxt.text.startswith('Research project') or (nxt.style.name=='List Bullet' and para.style.name!='List Bullet'):
    para.paragraph_format.keep_with_next=True
  for para in education_paras[:-1]:para.paragraph_format.keep_with_next=True
  if education_paras:education_paras[-1].paragraph_format.keep_with_next=False
  doc.core_properties.title='Research resume';doc.core_properties.subject='Application draft';doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
  dest=src.with_suffix('.docx');doc.save(dest);print(folder.name,'document created')
